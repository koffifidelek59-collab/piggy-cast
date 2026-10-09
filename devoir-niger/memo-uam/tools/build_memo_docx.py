"""Build the Word version of the Niger policy memo from the LaTeX source.

Pipeline: LaTeX (pre-processed for pandoc) -> pandoc + Lua filter -> .docx
          -> python-docx post-processing (cover page, header/footer, tables, fonts).
Equations are converted by pandoc into native, editable Word equations.

Usage: python3 build_docx.py <latex_dir> <output.docx>
"""
import copy
import os
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LATEX_DIR, OUT = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
TOOLS = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(LATEX_DIR, "niger_policy_memo.tex")

FONT = "Times New Roman"
GREEN = RGBColor(0x1E, 0x6B, 0x32)
RED = RGBColor(0xC0, 0x00, 0x00)
LEAF = RGBColor(0x6A, 0xA8, 0x4F)


# ------------------------------------------------------------------ 1. pre-process LaTeX
def simple_colspec(spec):
    spec = spec.replace("@{}", "")
    # expand *{n}{x}
    spec = re.sub(r"\*\{(\d+)\}\{(\w)\}", lambda m: m.group(2) * int(m.group(1)), spec)
    # drop >{...} groups (balanced braces)
    out, i = "", 0
    while i < len(spec):
        if spec.startswith(">{", i):
            depth, i = 0, i + 1
            while i < len(spec):
                depth += spec[i] == "{"
                depth -= spec[i] == "}"
                i += 1
                if depth == 0:
                    break
            continue
        out += spec[i]
        i += 1
    out = re.sub(r"L\{[^}]*\}", "l", out)
    out = re.sub(r"C\{[^}]*\}", "c", out)
    out = re.sub(r"[pm]\{[^}]*\}", "l", out)
    out = re.sub(r"[XY]", "l", out)
    return re.sub(r"[^lcr]", "", out)


def preprocess(tex):
    tex = re.sub(r"\\begin\{titlepage\}.*?\\end\{titlepage\}", "", tex, flags=re.S)
    for tok in ("\\newif\\ifcover", "\\covertrue", "\\ifcover", "\\setcounter{page}{1}"):
        tex = tex.replace(tok, "")
    tex = re.sub(r"^\\fi\s*$", "", tex, flags=re.M)
    # boxes -> quote blocks (styled as "Key Box" by the Lua filter)
    tex = re.sub(r"\\begin\{graybox\}\{([^}]*)\}", r"\\begin{quote}\n\\textbf{\1}\n", tex)
    tex = tex.replace("\\end{graybox}", "\\end{quote}")
    # page breaks via a marker paragraph
    tex = tex.replace("\\clearpage", "\n\nPAGEBREAKMARKER\n\n")
    # table cell helpers
    tex = re.sub(r"\\rng\{([^}]*)\}", r" [\1]", tex)
    tex = re.sub(r"\{\\footnotesize\\itshape ([^}]*)\}", r"\\textit{\1}", tex)
    tex = tex.replace("\\par\\vspace{2pt}", " LBRK ")
    tex = tex.replace("\\newline", " LBRK ")
    for a, b in (("$< -1.5$", "< \u22121.5"), ("$-1.5$", "\u22121.5"), ("$>$", ">"), ("$-$", "\u2212")):
        tex = tex.replace(a, b)

    def tabx(m):
        return "\\begin{tabular}{" + simple_colspec(m.group(1)) + "}"

    tex = re.sub(r"\\begin\{tabularx\}\{\\linewidth\}\{(.*?)\}\n", lambda m: tabx(m) + "\n", tex)
    tex = tex.replace("\\end{tabularx}", "\\end{tabular}")
    tex = re.sub(r"\\begin\{tabular\}\{l\*\{8\}\{c\}\}", r"\\begin{tabular}{lcccccccc}", tex)
    tex = tex.replace("\\textdegree", "°").replace("\\texteuro", "€")
    tex = tex.replace(r"\newcommand{\degC}{\,° C}", r"\newcommand{\degC}{\,°C}")
    tex = tex.replace(r"\newcommand{\pts}[1]{\hfill{\normalsize\textbf{(#1)}}}",
                      r"\newcommand{\pts}[1]{\quad\textbf{(#1)}}")
    tex = re.sub(r"\\needspace\{[^}]*\}", "", tex)
    tex = tex.replace("flow_kasso.pdf", "flow_kasso.png")
    tex = re.sub(r"\\renewcommand\{\\arraystretch\}\{[\d.]+\}", "", tex)
    # simple inline chemistry and symbols -> Unicode (robust, matches body font)
    sub = str.maketrans("0123456789+-", "₀₁₂₃₄₅₆₇₈₉₊₋")
    sup = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")
    tex = re.sub(r"\$_\{?([0-9]+)\}?\$", lambda m: m.group(1).translate(sub), tex)
    tex = re.sub(r"\$\^\{?([0-9]*[+-])\}?\$", lambda m: m.group(1).translate(sup), tex)
    tex = re.sub(r"\$_x\$", "ₓ", tex)
    for a, b in ((r"$\times$", "×"), (r"$\approx$", "≈"), (r"$\sim$", "~"), (r"$\mu$", "µ"),
                 (r"$\gamma$", "γ"), (r"$<$", "<"), (r"$\rightarrow$", "→")):
        tex = tex.replace(a, b)
    # subject titles: bold first line, plain second line
    tex = re.sub(r"\{\\LARGE\\bfseries ([^}]*)\}\\\\\[3pt\]\n\{\\large ([^}]*)\}",
                 r"\\textbf{\1}\\\\ \2", tex)
    # reference list -> plain paragraphs
    def refs(m):
        items = [i.strip() for i in m.group(1).split("\\item") if i.strip()]
        return "\n\n".join(items)
    tex = re.sub(r"\\begin\{list\}\{\}\{[^\n]*\}\n(.*?)\\end\{list\}", refs, tex, flags=re.S)
    return tex


with open(SRC, encoding="utf-8") as f:
    tex = preprocess(f.read())

work = tempfile.mkdtemp()
pre = os.path.join(LATEX_DIR, "_pandoc_tmp.tex")
with open(pre, "w", encoding="utf-8") as f:
    f.write(tex)


# ------------------------------------------------------------------ 2. reference document
ref = os.path.join(work, "reference.docx")
with open(ref, "wb") as f:
    f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                           capture_output=True, check=True).stdout)
rd = Document(ref)


def set_font(style, size=None, bold=None, italic=None, color=None):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), FONT)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = color if color is not None else RGBColor(0, 0, 0)


styles = rd.styles


def S(name):
    for st in styles:
        if st.name == name:
            return st
    raise KeyError(name)

for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in [s.name for s in styles]:
        st = S(name)
        set_font(st, 10.5)
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(3), 1.0
for name, size in (("Heading 1", 12.5), ("Heading 2", 12), ("Heading 3", 11)):
    st = S(name)
    set_font(st, size, bold=True)
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(8), Pt(3)
    st.paragraph_format.keep_with_next = True
for name in ("Title", "Subtitle"):
    set_font(S(name), 26 if name == "Title" else 14, bold=name == "Title")
for name in ("Image Caption", "Table Caption", "Caption"):
    if name in [s.name for s in styles]:
        set_font(S(name), 10, italic=False)
        S(name).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

from docx.enum.style import WD_STYLE_TYPE  # noqa: E402

ans = styles.add_style("Answer", WD_STYLE_TYPE.CHARACTER)
ans.font.name, ans.font.bold, ans.font.color.rgb = FONT, True, GREEN

kb = styles.add_style("Key Box", WD_STYLE_TYPE.PARAGRAPH)
kb.base_style = S("Normal")
set_font(kb, 10.5)
kbt = styles.add_style("Key Box Title", WD_STYLE_TYPE.PARAGRAPH)
kbt.base_style = S("Normal")
set_font(kbt, 11, bold=True)
for st in (kb, kbt):
    ppr = st.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "EFEFEF" if st is kb else "E0E0E0")
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", "000000")):
        left.set(qn(k), v)
    bdr.append(left)
    ppr.append(bdr)
    ppr.append(shd)
    st.paragraph_format.left_indent = Cm(0.3)
    st.paragraph_format.space_after = Pt(2)
kbt.paragraph_format.space_before = Pt(8)

cen = styles.add_style("Centered", WD_STYLE_TYPE.PARAGRAPH)
cen.base_style = S("Normal")
set_font(cen, 15, bold=False)
cen.paragraph_format.space_after = Pt(4)
cen.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cen.paragraph_format.space_after = Pt(8)
en = styles.add_style("End Note", WD_STYLE_TYPE.PARAGRAPH)
en.base_style = S("Normal")
set_font(en, 11, italic=True)
en.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
en.paragraph_format.space_before = Pt(10)
rd.save(ref)


# ------------------------------------------------------------------ 3. pandoc
raw = os.path.join(work, "raw.docx")
subprocess.run(["pandoc", pre, "-f", "latex", "-t", "docx", "--reference-doc", ref,
                "--lua-filter", os.path.join(TOOLS, "docx_filter.lua"),
                "--resource-path", LATEX_DIR, "-o", raw], check=True, cwd=LATEX_DIR)
os.remove(pre)


# ------------------------------------------------------------------ 4. post-processing
doc = Document(raw)
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.9)
TEXT_W = sec.page_width - sec.left_margin - sec.right_margin


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def booktabs(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", "12"), ("bottom", "12"), ("left", None), ("right", None),
                     ("insideH", None), ("insideV", None)):
        el = OxmlElement(f"w:{edge}")
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), sz)
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)
    # full width
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "pct")
    tblW.set(qn("w:w"), "5000")
    # header row rule + bold
    first = table.rows[0]
    for cell in first.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:color"), "000000")
        tcb.append(b)
        tcPr.append(tcb)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = FONT


WIDTHS = {  # first header cell -> column widths (cm), text width 17.2 cm
    "Impact (CIE indicator)": [2.9, 1.6, 2.3, 2.3, 8.1],
}


def grid(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)



def set_widths(table, widths):
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)


for t in doc.tables:
    booktabs(t)
    key = t.rows[0].cells[0].text.strip()
    if key.startswith("To:"):
        grid(t)
        for cell in t._cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for b in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(b)
        set_widths(t, [8.6, 8.6])
        continue
    if key in WIDTHS and len(WIDTHS[key]) == len(t.columns):
        set_widths(t, WIDTHS[key])

# images: limit width
for shape in doc.inline_shapes:
    if shape.width > TEXT_W:
        ratio = TEXT_W / shape.width
        shape.width = int(shape.width * ratio)
        shape.height = int(shape.height * ratio)

# ---- header / footer (not on the cover page)
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hp.text = ""
hp.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
r = hp.add_run("Climate Solutions and Strategies | Policy memo")
r.bold, r.font.size, r.font.name = True, Pt(9.5), FONT
r = hp.add_run("\tIMP-EGH | System Analysis | 2026")
r.font.size, r.font.name = Pt(9.5), FONT
pPr = hp._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bot = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
    bot.set(qn(k), v)
pbdr.append(bot)
pPr.append(pbdr)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run()
    run.font.size, run.font.name = Pt(10), FONT
    if kind:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        run._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = txt
        run._r.append(it)

# ---- cover page (built at the end of the document, then moved to the top)
body = doc.element.body
n_before = len(body)


def para(text="", size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER,
         after=0, before=0, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic, r.font.name = Pt(size), bold, italic, FONT
        if color:
            r.font.color.rgb = color
    return p


def initials(p, words, size, color, bold=True):
    for i, w in enumerate(words):
        if i:
            p.add_run(" ").font.size = Pt(size)
        if w[0].isupper() and w not in ("in", "and"):
            r = p.add_run(w[0])
            r.font.color.rgb, r.bold, r.font.size, r.font.name = color, bold, Pt(size), FONT
            r = p.add_run(w[1:])
        else:
            r = p.add_run(w)
        r.bold, r.font.size, r.font.name = bold, Pt(size), FONT


logo_tbl = doc.add_table(rows=1, cols=3)
no_borders(logo_tbl)
logo_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
c0, c1, c2 = logo_tbl.rows[0].cells
set_widths(logo_tbl, [3.5, 10.2, 3.5])
c0.paragraphs[0].add_run().add_picture(os.path.join(LATEX_DIR, "logo_uam.png"), width=Cm(3.1))
c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
c2.paragraphs[0].add_run().add_picture(os.path.join(LATEX_DIR, "logo_impegh.png"), width=Cm(3.3))
p = c1.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
initials(p, "Abdou Moumouni University".split(), 14, RED)
p = c1.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
initials(p, "International Master Program in Energy and Green Hydrogen".split(), 9.5, LEAF)
p = c1.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("(IMP-EGH)")
r.bold, r.font.size, r.font.name = True, Pt(11), FONT
for cell in (c0, c1, c2):
    cell.vertical_alignment = 1

rule = para(after=0, before=6)
rpPr = rule._p.get_or_add_pPr()
rb = OxmlElement("w:pBdr")
rbot = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "1"), ("w:color", "A59F93")):
    rbot.set(qn(k), v)
rb.append(rbot)
rpPr.append(rb)

para(before=130)
para("Why Limiting Climate Change", 26, bold=True)
para("Matters for Niger", 26, bold=True, after=12)
para("Policy memo to the Government of Niger", 16, italic=True, after=6)
para("Climate Solutions and Strategies     Assignment 1     Country case study: Niger", 12, after=22)
para("Assessing national climate risks with the Climate Impact Explorer", 12)
para("and the national interest of limiting global warming", 12, after=140)

sig = doc.add_table(rows=1, cols=2)
no_borders(sig)
set_widths(sig, [8.6, 8.6])
left, right = sig.rows[0].cells
for cell, lines, al in ((left, ["Course", "Climate Solutions and Strategies", "Assignment 1, due 14 October 2026"],
                         WD_ALIGN_PARAGRAPH.LEFT),
                        (right, ["Prepared by", "KOUAME Koffi Fidèle", "IMP-EGH, option System Analysis"],
                         WD_ALIGN_PARAGRAPH.RIGHT)):
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = al
        p.paragraph_format.space_after = Pt(4 if i == 0 else 0)
        r = p.add_run(line)
        r.bold, r.font.size, r.font.name = i == 0, Pt(11), FONT

para("October 2026", 11, bold=True, before=50)
brk = para()
brk.add_run().add_break(WD_BREAK.PAGE)

# move the cover elements (everything added after n_before, except the final sectPr) to the top
sectPr = body.find(qn("w:sectPr"))
new_elems = [el for el in list(body)[n_before - 1:] if el is not sectPr and el.tag != qn("w:sectPr")]
for i, el in enumerate(new_elems):
    body.remove(el)
    body.insert(i, el)

# ---- core properties
cp = doc.core_properties
cp.title = "Why Limiting Climate Change Matters for Niger | Policy memo"
cp.author = "KOUAME Koffi Fidèle"
cp.subject = "Climate Solutions and Strategies, Assignment 1"

doc.save(OUT)
print("saved", OUT)
